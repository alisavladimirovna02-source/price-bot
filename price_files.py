"""Read supplier price files without guessing prices from product names or IDs."""

import csv
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from itertools import chain
from pathlib import Path
import re
from zipfile import ZipFile


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000
MAX_CELL_LENGTH = 4096
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".txt"}
SUPPORTED_FORMATS = "CSV, Excel (.xlsx) или TXT"

HEADERS = {
    "name": {"наименование товара", "название товара", "наименование", "название",
             "товар", "name", "product", "product name"},
    "price": {"цена", "стоимость", "price"},
    "country": {"страна товара", "страна", "регион", "country", "region"},
}


class PriceFileError(ValueError):
    """An actionable file validation error safe to show in Telegram."""


def check_file(filename, size=0):
    if Path(filename or "").suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise PriceFileError(f"Пришли прайс в формате {SUPPORTED_FORMATS}.")
    if size > MAX_FILE_BYTES:
        raise PriceFileError("Файл слишком большой. Максимальный размер — 10 МБ.")


def decode_text(content):
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ("utf-16",)
    else:
        encodings = ("utf-8-sig", "cp1251")
    for encoding in encodings:
        try:
            text = content.decode(encoding)
        except UnicodeError:
            continue
        if any(ord(char) < 32 and char not in "\r\n\t" for char in text):
            raise PriceFileError("Файл не похож на текстовый прайс. Сохрани его как CSV или TXT в UTF-8.")
        return text
    raise PriceFileError("Не удалось прочитать кодировку. Сохрани файл в UTF-8.")


def normalize_header(value):
    return " ".join(str(value or "").lstrip("\ufeff").casefold().split())


def parse_amount(value):
    text = re.sub(r"\s+", "", str(value))
    text = re.sub(r"(?:₽|руб\.?|rub)$", "", text, flags=re.I)
    if not re.fullmatch(r"\d+(?:[.,]\d{1,2})?", text):
        raise PriceFileError("цена должна быть неотрицательным числом, например 19990 или 19 990,50")
    try:
        amount = Decimal(text.replace(",", "."))
        if amount > Decimal("1000000000000"):
            raise PriceFileError("слишком большая цена")
        return int(amount) if amount == amount.to_integral_value() else float(amount)
    except InvalidOperation as error:
        raise PriceFileError("не удалось прочитать цену") from error


def read_table(rows, start_row=1):
    columns = None
    entries = []
    for row_number, row in enumerate(rows, start_row):
        if row_number > MAX_ROWS + 1:
            raise PriceFileError(f"Слишком много строк. Максимум — {MAX_ROWS} строк товаров.")
        if len(row) > 256:
            raise PriceFileError("В таблице слишком много колонок (максимум 256).")
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if columns is None:
            headers = [normalize_header(value) for value in row]
            columns = {}
            for field, aliases in HEADERS.items():
                matches = [i for i, header in enumerate(headers) if header in aliases]
                if len(matches) > 1:
                    raise PriceFileError("В таблице несколько колонок с названием, ценой или страной. Оставь по одной.")
                if matches:
                    columns[field] = matches[0]
            if not {"name", "price"} <= columns.keys():
                raise PriceFileError("В первой строке таблицы нужны колонки «Наименование товара» (или «Название») и «Цена». Колонка «Страна товара» необязательна.")
            continue

        if len(row) > len(headers):
            raise PriceFileError(f"Строка {row_number}: больше значений, чем колонок. Проверь разделители и кавычки в CSV.")
        values = {field: str(row[index]).strip() if index < len(row) and row[index] is not None else ""
                  for field, index in columns.items()}
        if any(len(value) > MAX_CELL_LENGTH for value in values.values()):
            raise PriceFileError(f"Строка {row_number}: слишком длинное значение в ячейке.")
        if any(value.startswith("=") for value in values.values()):
            raise PriceFileError(f"Строка {row_number}: замени формулы в названии, цене и стране их значениями.")
        if not values["name"]:
            raise PriceFileError(f"Строка {row_number}: не заполнено название товара.")
        try:
            price = parse_amount(values["price"])
        except PriceFileError as error:
            raise PriceFileError(f"Строка {row_number}: {error}.") from error
        entries.append({"name": " ".join(values["name"].split()), "price": price,
                        "country": values.get("country", "").upper()})

    if not entries:
        raise PriceFileError("В файле нет товаров. Добавь строки с названиями и ценами.")
    return entries


def read_csv(content):
    text = decode_text(content).lstrip("\r\n")
    first, _, rest = text.partition("\n")
    if re.fullmatch(r"sep=[;,\t]\r?", first, re.I):
        delimiter, text = first[4], rest
    else:
        try:
            delimiter = csv.Sniffer().sniff(first, delimiters=";,\t").delimiter
        except csv.Error as error:
            raise PriceFileError("Не удалось определить колонки CSV. Используй разделитель «;», запятую или табуляцию.") from error
    try:
        return read_table(csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True))
    except csv.Error as error:
        raise PriceFileError("Повреждённый CSV: проверь разделители и кавычки.") from error


def read_xlsx(content):
    try:
        import openpyxl
    except ImportError as error:
        raise PriceFileError("Для Excel нужно обновить зависимости бота. Пока можно прислать CSV или TXT.") from error
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise PriceFileError("Excel-файл слишком большой после распаковки. Пришли таблицу меньшего размера.")
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            # Explicitly require one populated sheet to avoid silently losing offers.
            entries = None
            for sheet in workbook.worksheets:
                if sheet.sheet_state != "visible":
                    continue
                if sheet.max_column and sheet.max_column > 256:
                    raise PriceFileError("В Excel слишком много колонок. Сохрани таблицу без лишних колонок.")
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise PriceFileError(f"В Excel слишком много строк (максимум {MAX_ROWS} строк товаров).")
                rows = iter(sheet.iter_rows(values_only=True))
                first = next(((number, row) for number, row in enumerate(rows, 1)
                              if any(v is not None and str(v).strip() for v in row)), None)
                if first is None:
                    continue
                if entries is not None:
                    raise PriceFileError("В Excel несколько заполненных листов. Пришли каждый прайс отдельным файлом.")
                number, row = first
                entries = read_table(chain((row,), rows), start_row=number)
            if not entries:
                raise PriceFileError("В Excel нет товаров на видимых листах.")
            return entries
        finally:
            workbook.close()
    except PriceFileError:
        raise
    except Exception as error:
        raise PriceFileError("Не удалось прочитать Excel. Открой файл и сохрани его заново в формате .xlsx.") from error


def read_price_file(filename, content):
    """Return text lines or structured entries for the same price queue."""
    check_file(filename, len(content))
    if not content:
        raise PriceFileError("Файл пустой. Пришли прайс с товарами и ценами.")
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        return read_csv(content)
    if extension == ".xlsx":
        return read_xlsx(content)
    lines = [line.strip() for line in decode_text(content).splitlines() if line.strip()]
    if not lines:
        raise PriceFileError("Файл пустой. Пришли прайс с товарами и ценами.")
    if len(lines) > MAX_ROWS or any(len(line) > MAX_CELL_LENGTH for line in lines):
        raise PriceFileError(f"TXT слишком большой: максимум {MAX_ROWS} строк по {MAX_CELL_LENGTH} символов.")
    return lines
